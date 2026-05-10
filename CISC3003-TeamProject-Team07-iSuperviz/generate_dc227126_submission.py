from __future__ import annotations

import os
import shutil
import textwrap
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptInches, Pt as PptPt


PROJECT_DIR = Path(__file__).resolve().parent
REPO_DIR = PROJECT_DIR.parent
OUT_DIR = PROJECT_DIR / "individual_submission_DC227126"
SCREEN_DIR = PROJECT_DIR / "3003-screenshots"
PROJECT_FOLDER_NAME = "CISC3003-TeamProject-Team07-iSuperviz"
DEPLOY_DIR = PROJECT_DIR if PROJECT_DIR.name == PROJECT_FOLDER_NAME else REPO_DIR / PROJECT_FOLDER_NAME
CODE_ZIP = REPO_DIR / "cisc3003-Team07-ProjectCode.zip"

TEAM = "Team 07"
STUDENT_ID = "DC227126"
STUDENT_NAME = "Si Tin Iek"
PROJECT_TITLE = "iSuperviz: Your AI Research Supervisor"
GITHUB_URL = "https://github.com/Mikey-SI/CISC3003-dc227126-2026GitHub"
DEPLOYED_URL = "https://mikey-si.github.io/CISC3003-dc227126-2026GitHub/CISC3003-TeamProject-Team07-iSuperviz/"
DEMO_ACCOUNT = "professor@um.edu.mo / demo1234"
REDEEM_CODE = "TEAM07-ABC"

PURPLE = RGBColor(114, 46, 209)
DARK = RGBColor(36, 30, 56)


SCREENSHOTS = [
    ("01_home.png", "Home and brand entry point"),
    ("02_team.png", "Team 07 roster and project identity"),
    ("03_login.png", "Login and sign-up entry"),
    ("12_forgot_password.png", "Forgotten-password recovery"),
    ("04_setting.png", "Research preference setting"),
    ("05_paper_list.png", "Paper tracking list"),
    ("06_paper_detail.png", "Paper detail, reflective notes and AI chat"),
    ("07_idea_graph.png", "Idea Graph peculiar service"),
    ("08_hallucination.png", "Hallucination Audit peculiar service"),
    ("09_shop.png", "Research store"),
    ("10_cart.png", "Shopping cart and checkout"),
    ("11_dashboard.png", "Dashboard, orders and history"),
]


def ensure_dirs() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    DEPLOY_DIR.mkdir(exist_ok=True)


def set_doc_defaults(doc: Document) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = PURPLE if style_name != "Title" else DARK


def add_title(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = PURPLE if level > 0 else DARK


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value


def add_image_to_doc(doc: Document, image_path: Path, caption: str, width: float = 5.9) -> None:
    if not image_path.exists():
        doc.add_paragraph(f"[Missing screenshot: {image_path.name}]")
        return
    doc.add_picture(str(image_path), width=Inches(width))
    last = doc.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if cap.runs:
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(9)


def create_report() -> Path:
    doc = Document()
    set_doc_defaults(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(PROJECT_TITLE)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = PURPLE

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Individual Project Report").bold = True

    info = [
        f"Course: CISC3003 Web Programming",
        f"Team: {TEAM} - iSuperviz",
        f"Student: {STUDENT_NAME} ({STUDENT_ID})",
        "Pair: Pair 04",
        "Personal role: Frontend and backend development support, service integration, testing and demo preparation",
        "Submission deadline: 2026 May 10, 11:00 pm",
        f"GitHub Archive URL: {GITHUB_URL}",
        f"Deployed Testing URL: {DEPLOYED_URL}",
        f"Reviewer demo account: {DEMO_ACCOUNT}",
    ]
    for line in info:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    add_title(doc, "Abstract", 1)
    doc.add_paragraph(
        "This individual report explains my personal contribution to Team 07's iSuperviz project, "
        "a responsive full-stack web application designed as an AI research supervisor for university "
        "researchers. The team project combines React, TypeScript, Ant Design, React Flow, Node.js, "
        "Express, SQLite and a GitHub Pages-compatible mock API layer. My contribution focused on "
        "supporting the implementation of the website application together with Yang Xu, especially "
        "connecting frontend pages with backend service logic and checking that the user workflow could "
        "be demonstrated through the deployed URL. The work covered authentication-related flows, paper "
        "tracking, dashboard history, credit-aware actions, shopping cart behaviour, and the relationship "
        "between UI state and API responses. The individual evidence in this report follows the project "
        "requirements: context, goals, services and data requirements, tasks and tools, challenges, "
        "deliverables, division of work, individual contribution, success criteria, evaluation, learning, "
        "test cases, demo screenshots, references, AI declaration, installation instructions and incomplete "
        "items. The report also prepares the material for my individual digital story: a short presentation "
        "and screen demo explaining what I did, what services I helped integrate, how the project meets the "
        "CISC3003 requirements, and what I learned from building a deployable full-stack artifact."
    )

    add_title(doc, "Table of Contents", 1)
    toc_items = [
        "1. Project Context - Full-Stack Web Application with RWD and Deployed URL",
        "2. Personal Project Goals",
        "3. Services and Data Requirements I Worked With",
        "4. Personal Tasks and Tools",
        "5. Personal Challenges",
        "6. Deliverables Prepared for Submission",
        "7. Project Division of Work",
        "8. My Individual Contributions",
        "9. Criteria of Success",
        "10. Evaluation",
        "11. Individual and Team Learning",
        "12. Functional Test Cases",
        "13. Demo Walkthrough with Screenshots",
        "14. Installation, Deployment and Data Instructions",
        "15. Project Incomplete",
        "16. Citations and References",
        "17. Declaration of AI Usage",
        "18. Final Submission Checklist",
    ]
    add_numbered(doc, toc_items)
    doc.add_page_break()

    add_title(doc, "1. Project Context - Full-Stack Web Application with RWD and Deployed URL", 1)
    doc.add_paragraph(
        "iSuperviz is presented as a research-workflow platform rather than a simple static website. "
        "The frontend is a React single-page application with responsive layouts and project-specific "
        "visual pages. The source archive also includes a Node.js/Express/SQLite backend for authentication, "
        "paper data, ideas, chat messages, cart items, orders, search history and hallucination audit jobs. "
        "Because GitHub Pages cannot run a Node server, the deployed version uses a mock API adapter that "
        "mirrors the backend endpoints so the same workflow can still be tested publicly."
    )
    add_bullets(doc, [
        f"Deployed testing URL: {DEPLOYED_URL}",
        f"GitHub archive URL: {GITHUB_URL}",
        f"Demo account: {DEMO_ACCOUNT}",
        f"Redeem code for testing credits: {REDEEM_CODE}",
    ])

    add_title(doc, "2. Personal Project Goals", 1)
    add_bullets(doc, [
        "Support the frontend and backend implementation work assigned to Pair 04.",
        "Help connect UI pages with API/service logic so the demo can show real user workflows.",
        "Verify that the required modules are visible: login, password reset, dashboard, search, history and cart.",
        "Make sure project evidence can be explained clearly in the individual digital story.",
        "Learn how a full-stack project can be adapted for GitHub Pages deployment with mock API parity.",
    ])

    add_title(doc, "3. Services and Data Requirements I Worked With", 1)
    add_table(doc, ["Service Area", "Data / API Evidence", "Purpose"], [
        ["Authentication", "users, captchas, /api/login, /api/sendCaptcha", "Sign up, login and verification flow."],
        ["Password recovery", "/api/forgot_password, /api/check_account, /api/reset_password", "Recover user accounts."],
        ["Paper tracking", "papers, paper_notes, /api/get_paper_info, /api/get_paper_notes", "Search and read papers."],
        ["Idea services", "ideas, idea_related_papers", "Save ideas and visualise related papers."],
        ["Search history", "search_history, /api/search_history", "Record user activity."],
        ["Cart and orders", "products, cart_items, orders, /api/cart/*", "Support shopping-cart requirements."],
        ["Credits", "users.credit, /api/consume_credit, /api/redemption", "Deduct and top up research credits."],
        ["Audit service", "hallucination_jobs, /api/hallucination_check", "Provide the peculiar AI-audit feature."],
    ])

    add_title(doc, "4. Personal Tasks and Tools", 1)
    add_table(doc, ["Task", "My Work", "Tools"], [
        ["Frontend support", "Helped verify and connect pages such as Paper, Dashboard and Cart with service state.", "React, TypeScript, Ant Design"],
        ["Backend support", "Worked with service contracts for auth, paper, cart, order and history endpoints.", "Node.js, Express, SQLite"],
        ["Integration", "Checked how login state, credits, cart changes and history records update across pages.", "Axios, PubSub, localStorage"],
        ["Responsive/demo QA", "Reviewed screenshots and user journey for the individual explanation.", "Browser testing, screenshots"],
        ["Documentation", "Prepared individual report, PPT and narration based on team evidence.", "Word, PowerPoint, GitHub"],
    ])

    add_title(doc, "5. Personal Challenges", 1)
    add_bullets(doc, [
        "Keeping frontend state consistent after backend-style actions such as checkout, credit deduction and history updates.",
        "Explaining full-stack work honestly when the public deployment uses a static mock API due to GitHub Pages limitations.",
        "Making dense research pages understandable in a short individual digital story.",
        "Balancing team-level project evidence with my own individual contribution.",
        "Checking that screenshots, URLs and demo credentials can support assessment without extra setup.",
    ])

    add_title(doc, "6. Deliverables Prepared for Submission", 1)
    add_table(doc, ["Deliverable", "File / Evidence", "Status"], [
        ["Individual report", "CISC3003_Team07_DC227126_SiTinIek_Individual_Report.docx", "Prepared"],
        ["Individual PPT", "CISC3003_Team07_DC227126_SiTinIek_Individual_PPT.pptx", "Prepared"],
        ["Individual digital-story script", "CISC3003_Team07_DC227126_SiTinIek_Individual_Digital_Story_Script.docx", "Prepared"],
        ["Team proposal, report and PPT", "iSuperviz_Team07_Proposal (1)(1).docx, final report PDF and CISC3003_Team07_modified(1).pptx", "Prepared by team"],
        ["Coding archive", "cisc3003-Team07-ProjectCode.zip", "Generated"],
        ["Deployment folder", "CISC3003-TeamProject-Team07-iSuperviz", "Prepared for GitHub Pages path"],
    ])

    add_title(doc, "7. Project Division of Work", 1)
    add_table(doc, ["Member", "Pair", "Assigned work for final submission"], [
        ["Yang Xu", "Pair 08", "Overall lead, integration, deployment, report and major implementation coordination."],
        ["Jiang Xingyu", "Pair 08", "Public forum communication and written forum responses."],
        ["Huang Sofia", "Pair 12", "Project PPT materials, visual presentation structure and feature explanation."],
        ["Fan Zou Chen", "Pair 12", "Project PPT materials and presentation flow."],
        [STUDENT_NAME, "Pair 04", "Frontend and backend development support with service integration."],
        ["Ma Iat Tim", "Pair 04", "Citation/reference materials and APA reference organisation."],
    ])

    add_title(doc, "8. My Individual Contributions", 1)
    doc.add_paragraph(
        "My work was mainly implementation support and integration. I worked with Yang Xu on the website "
        "application so that the frontend pages were not isolated screens but connected to service logic. "
        "This included checking how authenticated user state is stored, how credits are shown after actions, "
        "how the cart and checkout update the dashboard, how search history is recorded, and how paper-related "
        "actions move the user toward reflective notes and the idea graph."
    )
    add_bullets(doc, [
        "Helped support frontend pages related to user workflow: Login, Forgot Password, Paper, Cart and Dashboard.",
        "Helped connect UI behaviour with backend/mock endpoints through Axios requests and shared response structures.",
        "Checked service integration points including login state, cart state, order history, search history and credit updates.",
        "Supported demo preparation by reviewing screenshots and organising the individual explanation of the deployed workflow.",
        "Contributed to the final project evidence by aligning personal contribution wording with the team report.",
    ])

    add_title(doc, "9. Criteria of Success", 1)
    add_bullets(doc, [
        "The user can open the deployed URL and enter the project without local setup.",
        "The demo account can access authenticated pages and show dashboard services.",
        "Paper search, paper detail, idea graph, hallucination audit, shop, cart and order history are demonstrable.",
        "Credits change visibly after research actions or checkout/redemption.",
        "Screenshots and narration can explain both team-level requirements and my individual role.",
    ])

    add_title(doc, "10. Evaluation", 1)
    add_table(doc, ["Evaluation Area", "Evidence", "Self-Evaluation"], [
        ["Functional requirements", "Login, reset, dashboard, search, history, cart and peculiar services visible.", "Meets course scope"],
        ["Integration", "Frontend requests match mock/backend service contracts.", "Effective for demo"],
        ["Responsive design", "Ant Design grids and screenshot walkthrough support different pages.", "Meets requirement"],
        ["Deployment", "GitHub Pages static bundle uses mock-first API parity.", "Testable"],
        ["Documentation", "Individual report, PPT and script explain personal contribution.", "Ready for submission"],
    ])

    add_title(doc, "11. Individual and Team Learning", 1)
    doc.add_paragraph(
        "Individually, I learned that frontend and backend development must share clear service contracts. "
        "A page is only convincing when UI state, API response, user feedback and stored data agree with each other. "
        "For example, cart checkout is not just a button: it must create an order, update credits, clear the cart, "
        "publish state changes and appear later in the dashboard. As a team, we learned that full-stack delivery "
        "also requires deployment evidence, screenshots, forum responses, references, AI-use transparency and "
        "individual accountability."
    )

    add_title(doc, "12. Functional Test Cases", 1)
    add_table(doc, ["Test Case", "Requirement Covered", "Expected / Verified Result"], [
        ["Login flow", "User login", "Demo account logs in and redirects to authenticated pages."],
        ["Forgot password", "Password recovery", "Account checking and reset form are visible."],
        ["Paper action", "Search and research workflow", "Explore/download actions deduct credits and open paper flow."],
        ["Cart checkout", "Shopping-cart service", "Cart accepts items, checkout creates order and updates credits."],
        ["Dashboard history", "History-related service", "Search/order/activity records appear in dashboard."],
        ["Idea graph", "Peculiar service", "Saved ideas and related papers appear as graph nodes."],
        ["Responsive view", "RWD requirement", "Main pages remain usable at desktop/tablet/mobile widths."],
    ])

    add_title(doc, "13. Demo Walkthrough with Screenshots", 1)
    doc.add_paragraph(
        "The following screenshots are used in my individual digital story. They follow the user journey from "
        "entry page to authentication, research workflow, peculiar services, e-commerce workflow and dashboard evidence."
    )
    for index, (file_name, caption) in enumerate(SCREENSHOTS, start=1):
        add_image_to_doc(doc, SCREEN_DIR / file_name, f"Figure {index}. {caption}.")

    add_title(doc, "14. Installation, Deployment and Data Instructions", 1)
    add_numbered(doc, [
        "Install Node.js 18/20/22 and npm.",
        f"Open the project folder {PROJECT_FOLDER_NAME}.",
        "Run npm install in the root folder.",
        "Run npm install --prefix server for the Express backend dependencies.",
        "Start the backend with npm start inside server.",
        "Start the frontend with npm start in the project root.",
        f"Open the deployed URL or local URL and login with {DEMO_ACCOUNT}.",
        "For GitHub Pages, build with CI=false and REACT_APP_FORCE_MOCK=1, then publish the build files under CISC3003-TeamProject-Team07-iSuperviz.",
    ])
    doc.add_paragraph(
        "Data usage: the backend uses SQLite. Seed/mock data includes products, paper catalogue, demo account, "
        "team roster, screenshots and deterministic audit outputs. No sensitive real user data is needed for assessment."
    )

    add_title(doc, "15. Project Incomplete", 1)
    add_bullets(doc, [
        "The public deployment uses a mock backend because GitHub Pages cannot run Node.js.",
        "Hallucination Audit uses deterministic demo logic instead of a paid production detector.",
        "Email verification falls back to console/demo mode unless SMTP credentials are configured.",
        "The paper catalogue is curated/seeded rather than a live arXiv crawler.",
        "The individual and team digital story videos must still be recorded as MP4 files from the PPT and screen demo.",
    ])

    add_title(doc, "16. Citations and References", 1)
    references = [
        "Ant Design Team. (n.d.). Ant Design - A design system for enterprise-level products. https://ant.design/",
        "Cornell University. (n.d.). arXiv.org e-Print archive. https://arxiv.org/",
        "Express.js. (n.d.). Express - Fast, unopinionated, minimalist web framework for Node.js. https://expressjs.com/",
        "GitHub, Inc. (n.d.). GitHub Pages documentation. https://docs.github.com/en/pages",
        "Meta Open Source. (n.d.). React documentation. https://react.dev/",
        "Microsoft. (n.d.). TypeScript documentation. https://www.typescriptlang.org/docs/",
        "Node.js Foundation. (n.d.). Node.js documentation. https://nodejs.org/en/docs/",
        "Nodemailer. (n.d.). Nodemailer - Send emails from Node.js. https://nodemailer.com/",
        "React Flow. (n.d.). React Flow - Build node-based UIs with React. https://reactflow.dev/",
        "SQLite Consortium. (n.d.). SQLite documentation. https://www.sqlite.org/docs.html",
        "World Wide Web Consortium. (2018). Web Content Accessibility Guidelines (WCAG) 2.1. https://www.w3.org/TR/WCAG21/",
    ]
    for ref in references:
        doc.add_paragraph(ref)

    add_title(doc, "17. Declaration of AI Usage", 1)
    doc.add_paragraph(
        "AI assistants, including ChatGPT, Claude and Cursor, were used as assistive tools for wording, "
        "debugging support, checklist mapping, report/PPT organisation and review of missing requirements. "
        "AI was used under human supervision. Final code, screenshots, design decisions and written content "
        "were checked and edited by the student/team. AI was not used as a substitute for learning, judgement "
        "or responsibility, and no sensitive personal data was provided to AI services."
    )

    add_title(doc, "18. Final Submission Checklist", 1)
    add_bullets(doc, [
        "Submit team proposal + team report + team PPT to the teamwork project link.",
        "Submit cisc3003-Team07-ProjectCode.zip to the coding archive link.",
        "Submit team digital story MP4/ZIP to the team digital story link.",
        "Submit individual digital story MP4/ZIP, created from the individual PPT and script, to the individual digital story link.",
        "Post GitHub archive URL, deployed testing URL, team digital story URL and individual digital story URL in the collaborative forum threads.",
    ])

    out = OUT_DIR / "CISC3003_Team07_DC227126_SiTinIek_Individual_Report.docx"
    doc.save(out)
    return out


SLIDES = [
    {
        "title": "iSuperviz Individual Digital Story",
        "bullets": [
            "Si Tin Iek (DC227126), Pair 04, Team 07",
            "Personal focus: frontend/backend support and service integration",
            "Project: AI research supervisor full-stack web application",
        ],
        "images": ["01_home.png"],
    },
    {
        "title": "Course Requirements and My Role",
        "bullets": [
            "RWD, full-stack services, dashboard, search, history and cart",
            "Worked with Yang Xu on implementation support",
            "Checked that UI pages connect to service logic",
        ],
        "images": ["02_team.png"],
    },
    {
        "title": "Full-Stack Architecture",
        "bullets": [
            "React + TypeScript frontend",
            "Express + SQLite backend in source archive",
            "Mock API parity for GitHub Pages deployment",
        ],
        "images": ["11_dashboard.png"],
    },
    {
        "title": "Authentication and Password Recovery",
        "bullets": [
            "Login flow uses demo account for assessment",
            "Forgot-password page supports account checking and reset",
            "Backend defines captcha/email verification logic",
        ],
        "images": ["03_login.png", "12_forgot_password.png"],
    },
    {
        "title": "Paper Tracking Workflow",
        "bullets": [
            "Paper list supports search, favorites and credit-aware actions",
            "Paper detail shows notes, ideas and AI chat",
            "Credits and state changes make the workflow testable",
        ],
        "images": ["05_paper_list.png", "06_paper_detail.png"],
    },
    {
        "title": "Peculiar Research Services",
        "bullets": [
            "Idea Graph visualises saved ideas and related papers",
            "Hallucination Audit demonstrates AI-review service",
            "These features extend the project beyond basic login/cart pages",
        ],
        "images": ["07_idea_graph.png", "08_hallucination.png"],
    },
    {
        "title": "Shopping Cart and Dashboard",
        "bullets": [
            "Shop and cart implement e-commerce requirements",
            "Checkout updates order records and credits",
            "Dashboard brings profile, cart, history and orders together",
        ],
        "images": ["09_shop.png", "10_cart.png", "11_dashboard.png"],
    },
    {
        "title": "Testing, Deployment and Evidence",
        "bullets": [
            "Team tests covered mock API parity and responsive rendering",
            "GitHub Pages provides public testing URL",
            "Screenshots document the full demo walkthrough",
        ],
        "images": ["01_home.png"],
    },
    {
        "title": "Personal Learning and Challenges",
        "bullets": [
            "Learned the importance of shared API contracts",
            "Managed state consistency across credits, cart and history",
            "Practiced explaining full-stack work clearly and honestly",
        ],
        "images": ["04_setting.png"],
    },
    {
        "title": "Submission Wrap-Up",
        "bullets": [
            "Prepared individual report, PPT and narration script",
            "Coding archive named cisc3003-Team07-ProjectCode.zip",
            "Digital story should be recorded as 10-15 minute MP4",
        ],
        "images": ["01_home.png"],
    },
]


SCRIPT = [
    (
        "Slide 1",
        "Hello, my name is Si Tin Iek, student ID DC227126, from Pair 04 in Team 07. "
        "In this individual digital story, I will explain my contribution to iSuperviz, our AI research supervisor web application. "
        "My main role was supporting frontend and backend development with Yang Xu, especially helping the website pages connect with service logic. "
        "I will focus on what I worked on, how the services fit the course requirements, and what I learned from turning a project idea into a deployed web artifact."
    ),
    (
        "Slide 2",
        "The CISC3003 project required a responsive full-stack web application with user login, password reset, dashboard, search, history, shopping cart, peculiar services and a deployed URL. "
        "Inside the team, I focused on implementation support. I checked that the screens were not only visual designs, but could follow the required user workflow."
        " This means that when a user logs in, searches, spends credits, checks the cart, or opens the dashboard, the interface should reflect a clear service behind it."
    ),
    (
        "Slide 3",
        "The project architecture has a React and TypeScript frontend, an Express and SQLite backend, and a mock API layer for GitHub Pages. "
        "This was important because GitHub Pages cannot run a Node server. I learned how the frontend can still call the same style of API responses during the public deployed demo."
        " The local backend remains in the coding archive, while the public version is designed for convenient assessment through one URL."
    ),
    (
        "Slide 4",
        "For the account workflow, the login page allows assessors to use the demo account, while the forgot-password page demonstrates account checking and reset logic. "
        "I reviewed how the frontend stores login information, redirects unauthenticated users, and depends on backend or mock responses for a complete authentication flow."
        " This part also helped me understand why small validation details matter, because account services are the first gate before the dashboard and research functions."
    ),
    (
        "Slide 5",
        "The paper tracking workflow is the core research feature. Users can search papers, open paper details, read reflective notes, ask questions and use credit-based actions. "
        "My integration focus was to understand how paper actions affect user credits and how frontend state should update after a service call."
        " For example, Explore and Download PDF should not feel like static buttons; they should behave like real research services with visible cost and feedback."
    ),
    (
        "Slide 6",
        "The peculiar services show our own project idea beyond ordinary course requirements. The Idea Graph connects saved research ideas with related papers, and Hallucination Audit gives a structured AI-review result. "
        "These features helped the project feel like a research supervisor instead of only a shopping or login system."
        " When preparing my explanation, I used these two pages as evidence that the project has a domain-specific purpose."
    ),
    (
        "Slide 7",
        "The shop, cart and dashboard cover the e-commerce and history-related requirements. A user can add credit products, check out, receive updated credits, and later see orders and activity history in the dashboard. "
        "This part taught me that one action must update several parts of the application consistently."
        " If checkout succeeds but the dashboard does not change, the workflow is not convincing, so integration checking is very important."
    ),
    (
        "Slide 8",
        "For testing and deployment evidence, the team used unit tests for mock API parity and responsive rendering, plus manual browser checks and screenshots. "
        "The deployed GitHub Pages version uses a mock-first bundle so assessors can test the main functions through one public URL without installing the backend."
        " The screenshots in my report and PPT are arranged as a walkthrough so that the video can show evidence quickly and clearly."
    ),
    (
        "Slide 9",
        "My biggest learning was that frontend and backend must agree on service contracts. If the API response, local storage, page state and dashboard display do not match, the user experience breaks. "
        "I also learned how to explain limitations honestly, especially the difference between local Express backend and GitHub Pages mock deployment."
        " This project made me more aware that documentation, testing and deployment are part of engineering work, not separate tasks after coding."
    ),
    (
        "Slide 10",
        "To conclude, my individual contribution was supporting implementation and service integration for the iSuperviz website. "
        "For submission, I prepared this individual PPT, report and narration script, while the team code archive is named cisc3003-Team07-ProjectCode.zip. "
        "The final individual digital story should record this narration with the PPT and a short deployed website demo. "
        "During recording, I will open the deployed URL, login with the demo account, and briefly show the paper, cart and dashboard pages as proof of my contribution."
    ),
]


def add_ppt_textbox(slide, left, top, width, height, text, font_size=22, bold=False, color=PptRGBColor(36, 30, 56)):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = PptPt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def add_ppt_bullets(slide, bullets: list[str], left, top, width, height) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = PptPt(18)
        p.font.color.rgb = PptRGBColor(45, 39, 61)


def add_ppt_image(slide, image_path: Path, left, top, width, height) -> None:
    if not image_path.exists():
        return
    with Image.open(image_path) as img:
        iw, ih = img.size
    target_ratio = width / height
    image_ratio = iw / ih
    if image_ratio > target_ratio:
        pic_h = height
        pic_w = height * image_ratio
    else:
        pic_w = width
        pic_h = width / image_ratio
    pic_left = left + (width - pic_w) / 2
    pic_top = top + (height - pic_h) / 2
    slide.shapes.add_picture(str(image_path), pic_left, pic_top, width=pic_w, height=pic_h)


def create_ppt() -> Path:
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    blank = prs.slide_layouts[6]

    for idx, data in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = PptRGBColor(248, 245, 255)

        slide.shapes.add_shape(1, PptInches(0), PptInches(0), PptInches(13.333), PptInches(0.25)).fill.solid()
        slide.shapes[-1].fill.fore_color.rgb = PptRGBColor(114, 46, 209)
        slide.shapes[-1].line.fill.background()

        add_ppt_textbox(slide, PptInches(0.55), PptInches(0.45), PptInches(9.5), PptInches(0.55), data["title"], 28, True, PptRGBColor(114, 46, 209))
        add_ppt_textbox(slide, PptInches(11.7), PptInches(0.5), PptInches(1.1), PptInches(0.35), f"{idx}/10", 14, False, PptRGBColor(100, 91, 130))
        add_ppt_bullets(slide, data["bullets"], PptInches(0.65), PptInches(1.25), PptInches(5.1), PptInches(5.6))

        images = data.get("images", [])
        if len(images) == 1:
            add_ppt_image(slide, SCREEN_DIR / images[0], PptInches(6.05), PptInches(1.2), PptInches(6.7), PptInches(5.65))
        elif len(images) == 2:
            add_ppt_image(slide, SCREEN_DIR / images[0], PptInches(5.9), PptInches(1.25), PptInches(3.35), PptInches(5.4))
            add_ppt_image(slide, SCREEN_DIR / images[1], PptInches(9.35), PptInches(1.25), PptInches(3.35), PptInches(5.4))
        else:
            add_ppt_image(slide, SCREEN_DIR / images[0], PptInches(5.85), PptInches(1.2), PptInches(3.0), PptInches(5.4))
            add_ppt_image(slide, SCREEN_DIR / images[1], PptInches(8.75), PptInches(1.2), PptInches(2.2), PptInches(2.55))
            add_ppt_image(slide, SCREEN_DIR / images[2], PptInches(8.75), PptInches(4.05), PptInches(3.9), PptInches(2.8))

        footer = slide.shapes.add_textbox(PptInches(0.65), PptInches(7.03), PptInches(12.0), PptInches(0.25))
        p = footer.text_frame.paragraphs[0]
        p.text = "CISC3003 Team 07 - iSuperviz - Si Tin Iek (DC227126)"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = PptPt(10)
        p.font.color.rgb = PptRGBColor(100, 91, 130)

    out = OUT_DIR / "CISC3003_Team07_DC227126_SiTinIek_Individual_PPT.pptx"
    prs.save(out)
    return out


def create_script_docs() -> tuple[Path, Path]:
    plain = OUT_DIR / "CISC3003_Team07_DC227126_SiTinIek_Individual_Digital_Story_Script.txt"
    docx = OUT_DIR / "CISC3003_Team07_DC227126_SiTinIek_Individual_Digital_Story_Script.docx"

    text = ["Individual Digital Story Script", f"{STUDENT_NAME} ({STUDENT_ID}), {TEAM}", ""]
    for slide, paragraph in SCRIPT:
        text.append(f"{slide}:")
        text.append(textwrap.fill(paragraph, width=100))
        text.append("")
    plain.write_text("\n".join(text), encoding="utf-8")

    doc = Document()
    set_doc_defaults(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Individual Digital Story Script")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = PURPLE
    subtitle = doc.add_paragraph(f"{STUDENT_NAME} ({STUDENT_ID}), {TEAM}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Approximate length: around 700 English words. Use this while recording voice narration over the PPT and deployed website demo.")
    for slide, paragraph in SCRIPT:
        add_title(doc, slide, 2)
        doc.add_paragraph(paragraph)
    doc.save(docx)
    return plain, docx


def create_checklist() -> Path:
    checklist = OUT_DIR / "CISC3003_Team07_DC227126_Submission_Checklist.txt"
    content = f"""CISC3003 Team07 DC227126 Submission Checklist

Student: {STUDENT_NAME} ({STUDENT_ID})
Team: {TEAM} - iSuperviz

1. Teamwork Project Assignment (Proposal + Report + PPT)
Submit:
- iSuperviz_Team07_Proposal (1)(1).docx
- iSuperviz_Team07_Project_Report_Final(1).pdf
- CISC3003_Team07_modified(1).pptx

2. Teamwork Project Assignment (Coding Archive)
Submit:
- cisc3003-Team07-ProjectCode.zip
Also push the project to:
- {GITHUB_URL}

3. Collaborative Forum Threads
Create/post:
- Our Project GitHub Link: {GITHUB_URL}
- Our Project Deployed URL for Testing: {DEPLOYED_URL}
- My Team Digital Story for Project Assignment: paste the team video link after upload
- My Individual Digital Story for Project Assignment: paste your own video link after upload

4. Team Digital Story
Submit:
- Team MP4 or ZIP, 10-15 minutes, under 200 MB
- Must include every member's participation

5. Individual Digital Story
Use:
- CISC3003_Team07_DC227126_SiTinIek_Individual_PPT.pptx
- CISC3003_Team07_DC227126_SiTinIek_Individual_Digital_Story_Script.docx
Record:
- PPT voice narration plus deployed website screen demo
- 10-15 minutes, under 200 MB
Preferred ZIP name:
- cisc3003-Team07-DC227126-ProjectDigiStory.zip

Demo login:
- {DEMO_ACCOUNT}
Redeem code:
- {REDEEM_CODE}
"""
    checklist.write_text(content, encoding="utf-8")
    return checklist


def prepare_deploy_folder() -> None:
    build_dir = PROJECT_DIR / "build"
    if not build_dir.exists():
        return
    shutil.copytree(build_dir, DEPLOY_DIR, dirs_exist_ok=True)
    (DEPLOY_DIR / ".nojekyll").write_text("", encoding="utf-8")
    index = DEPLOY_DIR / "index.html"
    if index.exists():
        shutil.copy2(index, DEPLOY_DIR / "404.html")


def should_skip_zip(path: Path) -> bool:
    parts = set(path.parts)
    skip_dirs = {".git", "node_modules", ".cache", "__pycache__"}
    if parts.intersection(skip_dirs):
        return True
    if path.name in {CODE_ZIP.name}:
        return True
    if path.suffix.lower() in {".pyc", ".log"}:
        return True
    return False


def create_code_zip() -> Path:
    if CODE_ZIP.exists():
        CODE_ZIP.unlink()
    roots = list(dict.fromkeys([PROJECT_DIR, DEPLOY_DIR]))
    with zipfile.ZipFile(CODE_ZIP, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as zf:
        for root in roots:
            if not root.exists():
                continue
            for file_path in root.rglob("*"):
                if file_path.is_dir() or should_skip_zip(file_path):
                    continue
                arcname = file_path.relative_to(REPO_DIR)
                zf.write(file_path, arcname.as_posix())
    return CODE_ZIP


def main() -> None:
    ensure_dirs()
    report = create_report()
    ppt = create_ppt()
    script_txt, script_docx = create_script_docs()
    checklist = create_checklist()
    prepare_deploy_folder()
    code_zip = create_code_zip()
    print("Generated files:")
    for path in [report, ppt, script_docx, script_txt, checklist, code_zip, DEPLOY_DIR]:
        print(f"- {path}")


if __name__ == "__main__":
    main()
